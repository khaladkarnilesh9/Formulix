from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from apps.core.services import transition
from apps.formulations.models import Formulation
from apps.formulations.spreadsheet_columns import COLUMNS
from apps.imports.formulation_excel import commit_workbook, preview_workbook
from apps.imports.models import ImportBatch


class FormulationExcelTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.operator = get_user_model().objects.create_superuser(username="excel-operator", password=None)
        cls.other = get_user_model().objects.create_user(username="excel-user", password=None, is_staff=True)

    def upload(self, invalid=False, formula=False):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Master Database"
        sheet.append(list(COLUMNS))
        row = ["Example company", "" if invalid else "Example SC", "Insecticide", "SC", "Ingredient A 20% + Ingredient B 8% w/w",
               "Dispersant (CAS 37199-81-8)", "Water (CAS 7732-18-5)", "Glycol + gum", "1.08 - 1.12 g/mL",
               "6.5 - 7.5", "Source REG-1", "> 100 C", "400 - 600 cP", "Class III", "Light Blue", "Source Patent-1", "Armyworm"]
        if formula:
            row[4] = "=1+1"
        sheet.append(row)
        copy = workbook.create_sheet("Example company")
        copy.append(list(COLUMNS))
        copy.append(row)
        copy.append([*row[:1], "Second SC", *row[2:]])
        output = BytesIO()
        workbook.save(output)
        return SimpleUploadedFile("sample.xlsx", output.getvalue())

    def test_preview_all_sheets_deduplicates_and_keeps_missing_master_rows(self):
        preview = preview_workbook(self.upload(), "India")
        self.assertEqual(preview["new_count"], 2)
        self.assertEqual(preview["duplicates"], 1)
        self.assertEqual(preview["rows"][0]["raw"]["pH Range"], "6.5 - 7.5")
        self.assertFalse(Formulation.objects.exists())

    def test_master_provenance_wins_when_manufacturer_sheet_is_first(self):
        workbook = load_workbook(BytesIO(self.upload().read()))
        workbook.move_sheet("Master Database", offset=1)
        data = BytesIO()
        workbook.save(data)
        preview = preview_workbook(SimpleUploadedFile("master-last.xlsx", data.getvalue()), "India")
        self.assertEqual(preview["duplicates"], 1)
        self.assertEqual(len(preview["rows"]), 2)
        self.assertEqual(preview["rows"][0]["sheet"], "Master Database")

    def test_import_is_draft_exact_audited_and_idempotent(self):
        preview = preview_workbook(self.upload(), "India")
        batch = commit_workbook(preview, self.operator)
        record = Formulation.objects.get(formulation_name="Example SC")
        self.assertEqual(record.publication_status, "Draft")
        self.assertEqual(record.verification_status, "Unverified")
        self.assertEqual(record.active_ingredient_text, preview["rows"][0]["raw"]["Active Ingredient & % (w/w)"])
        self.assertEqual(record.legacy_payload["spreadsheetRow"], preview["rows"][0]["raw"])
        self.assertEqual(batch.imported_rows, 2)
        self.assertEqual(batch.rows.filter(import_status="Imported").count(), 2)
        self.assertEqual(commit_workbook(preview, self.operator).pk, batch.pk)
        self.assertEqual(Formulation.objects.count(), 2)
        with self.assertRaises(ValidationError):
            commit_workbook({**preview, "market": "Europe"}, self.operator)

    def test_permissions_and_invalid_rows(self):
        preview = preview_workbook(self.upload(invalid=True), "Africa")
        with self.assertRaises(PermissionDenied):
            commit_workbook(preview, self.other)
        batch = commit_workbook(preview, self.operator)
        self.assertEqual(batch.rejected_rows, 1)
        self.assertEqual(batch.imported_rows, 1)
        self.assertTrue(batch.error_summary)

    def test_formulas_are_rejected_without_execution(self):
        preview = preview_workbook(self.upload(formula=True), "China")
        self.assertEqual(preview["error_count"], 2)
        self.assertEqual(preview["new_count"], 0)

    def test_published_api_contains_all_columns_and_properties(self):
        commit_workbook(preview_workbook(self.upload(), "United States"), self.operator)
        record = Formulation.objects.get(formulation_name="Example SC")
        url = f"/api/v1/formulations/{record.pk}/"
        self.assertEqual(self.client.get(url).status_code, 404)
        for status in ("In Review", "Approved", "Published"):
            record = transition(record, status, self.operator)
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(set(data["marketData"]), {"Market", *COLUMNS})
        self.assertEqual(data["properties"]["viscosity"], "400 - 600 cP")
        self.assertEqual(data["composition"], [])
        self.assertEqual(data["patentReferences"], [])

    def test_admin_upload_preview_then_confirm(self):
        self.client.force_login(self.operator)
        url = reverse("admin:formulations_formulation_upload_excel")
        self.assertEqual(self.client.get(url).status_code, 200)
        preview = self.client.post(url, {"market": "Europe", "workbook": self.upload()})
        self.assertEqual(preview.status_code, 200)
        self.assertContains(preview, "Confirm Draft Import")
        token = self.client.session["formulation_excel_preview"]["token"]
        self.assertFalse(ImportBatch.objects.exists())
        response = self.client.post(url, {"confirm": "1", "token": token})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Formulation.objects.count(), 2)
        self.assertNotIn("formulation_excel_preview", self.client.session)

    def test_unknown_market_and_wrong_file_rejected(self):
        with self.assertRaises(ValidationError):
            preview_workbook(self.upload(), "Unspecified")
        with self.assertRaises(ValidationError):
            preview_workbook(SimpleUploadedFile("file.csv", b"data"), "India")
