import hashlib
import json
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from django.core.exceptions import PermissionDenied, ValidationError
from django.db import transaction
from openpyxl import load_workbook

from apps.core.services import audit
from apps.formulations.models import Formulation
from apps.formulations.spreadsheet_columns import COLUMNS, MARKETS
from apps.imports.models import ImportBatch, ImportRow
from apps.taxonomy.models import IndustrySector

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10000


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def preview_workbook(upload, market):
    if market not in MARKETS:
        raise ValidationError("Choose the market explicitly; filenames are not evidence of jurisdiction.")
    filename = Path(upload.name).name
    if not filename.lower().endswith(".xlsx"):
        raise ValidationError("Upload an .xlsx workbook.")
    data = upload.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise ValidationError("Workbook exceeds the 10 MB upload limit.")
    try:
        with ZipFile(BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 80 * 1024 * 1024:
                raise ValidationError("Expanded workbook exceeds the safety limit.")
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, OSError, ValueError, KeyError) as error:
        raise ValidationError("Invalid Excel workbook.") from error
    rows, seen, duplicates, scanned, sheets = [], set(), 0, 0, []
    try:
        for sheet in sorted(workbook, key=lambda sheet: sheet.title.strip().casefold() != "master database"):
            iterator = sheet.iter_rows()
            headings = [cell_text(cell.value).strip() for cell in next(iterator, ())]
            if not set(COLUMNS).issubset(headings):
                continue
            if len(set(headings)) != len(headings):
                raise ValidationError(f"Duplicate column headings in {sheet.title}.")
            sheets.append(sheet.title)
            indexes = {heading: headings.index(heading) for heading in COLUMNS}
            for number, cells in enumerate(iterator, start=2):
                scanned += 1
                if scanned > MAX_ROWS * 3:
                    raise ValidationError("Workbook has too many rows.")
                raw = {heading: cell_text(cells[index].value) if index < len(cells) else "" for heading, index in indexes.items()}
                if not any(value.strip() for value in raw.values()):
                    continue
                fingerprint = hashlib.sha256(json.dumps([market, raw], sort_keys=True).encode()).hexdigest()
                if fingerprint in seen:
                    duplicates += 1
                    continue
                seen.add(fingerprint)
                errors = []
                if any(cell.data_type == "f" for cell in cells):
                    errors.append("Formula cells are not accepted; upload literal values.")
                name = raw["Product Name / Brand Name"]
                if not name.strip() or len(name) > 300:
                    errors.append("Product name is required and must not exceed 300 characters.")
                if not raw["Company Name"].strip():
                    errors.append("Company name is required.")
                if len(raw["Formulation Type"]) > 300:
                    errors.append("Formulation type exceeds 300 characters.")
                if max(map(len, raw.values())) > 20000:
                    errors.append("Cell text exceeds the safety limit.")
                rows.append({"raw": raw, "key": fingerprint, "sheet": sheet.title, "row": number, "errors": errors})
                if len(rows) > MAX_ROWS:
                    raise ValidationError("Workbook exceeds 10,000 unique rows.")
    finally:
        workbook.close()
    if not sheets or not rows:
        raise ValidationError("No data found with the expected 17 spreadsheet columns.")
    existing = set(Formulation.objects.filter(import_key__in=seen).values_list("import_key", flat=True))
    for row in rows:
        row["exists"] = row["key"] in existing
    return {"filename": filename[:300], "market": market, "sha256": hashlib.sha256(data).hexdigest(),
            "rows": rows, "sheets": sheets, "duplicates": duplicates,
            "new_count": sum(not row["errors"] and not row["exists"] for row in rows),
            "error_count": sum(bool(row["errors"]) for row in rows), "existing_count": len(existing)}


@transaction.atomic
def commit_workbook(preview, actor):
    if not actor.is_staff or not actor.has_perms(("formulations.add_formulation", "imports.add_importbatch", "imports.add_importrow")):
        raise PermissionDenied("Formulation and import creation permissions are required.")
    batch, created = ImportBatch.objects.get_or_create(
        uploaded_by=actor, import_type="Market formulations XLSX", sha256=preview["sha256"],
        defaults={"filename": preview["filename"], "status": "Importing", "total_rows": len(preview["rows"]),
                  "field_mapping": {"market": preview["market"], "columns": COLUMNS}})
    if not created:
        if batch.field_mapping.get("market") != preview["market"]:
            raise ValidationError("This workbook was already imported under another market. Review that batch first.")
        return batch
    sector, _ = IndustrySector.objects.get_or_create(name="Agrochemical", defaults={"slug": "agrochemical"})
    errors = []
    for index, row in enumerate(preview["rows"], start=1):
        entry = ImportRow(batch=batch, row_number=index, raw_data=row["raw"],
                          normalized_data={"market": preview["market"], "sheet": row["sheet"], "sourceRow": row["row"]})
        record = Formulation.objects.filter(import_key=row["key"]).first()
        if row["errors"]:
            entry.import_status, entry.validation_errors = "Rejected", row["errors"]
            batch.rejected_rows += 1
            errors.append({"sheet": row["sheet"], "row": row["row"], "errors": row["errors"]})
        elif record:
            entry.import_status, entry.formulation = "Skipped", record
            batch.valid_rows += 1
        else:
            values = {field: row["raw"][heading] for heading, field in COLUMNS.items()}
            record = Formulation(**values, market_region=preview["market"], import_key=row["key"],
                                 industry_sector=sector, source_title=preview["filename"], source_type="Uploaded spreadsheet",
                                 source_reference=f'{preview["filename"]} / {row["sheet"]} / row {row["row"]}',
                                 source_sheet=row["sheet"], source_row=row["row"], created_by=actor,
                                 intended_application=values["target_pests_text"], target_industry="Agrochemical",
                                 verification_label="Imported unverified", legacy_payload={"spreadsheetRow": row["raw"]})
            record.full_clean()
            record.save()
            audit(actor, record, "Imported", record.source_reference)
            entry.import_status, entry.formulation = "Imported", record
            batch.valid_rows += 1
            batch.imported_rows += 1
        entry.full_clean()
        entry.save()
    batch.status = "Completed with errors" if batch.rejected_rows else "Completed"
    batch.error_summary = errors
    batch.save()
    return batch
